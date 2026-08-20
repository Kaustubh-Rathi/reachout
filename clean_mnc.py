#!/usr/bin/env python3
"""Clean the MNC sheet from Reachout.xlsx and export a streamlined Excel file.

Extracts strictly 4 columns:
  - Company Name
  - Person Name
  - Phone Number
  - Email

Delegates all core domain normalization and deduplication to contact_ingestion.py.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from contact_ingestion import (
    Contact,
    clean_contacts,
    clean_text,
    find_default_workbook,
    read_mnc_rows,
)

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DEFAULT_INPUT = find_default_workbook()
DEFAULT_OUTPUT_XLSX = DATA_DIR / "MNC_cleaned.xlsx"


def save_to_excel(output_path: Path, contacts: List[Contact]) -> Path:
    """Save contacts to an Excel file with clean styling."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MNC_Cleaned"

    # Style definitions
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    headers = ["Company Name", "Person Name", "Phone Number", "Email"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for row_idx, c in enumerate(contacts, start=2):
        ws.append([c.company, c.name, f"+{c.phone}" if c.phone else "", c.email])

        use_alt = (row_idx % 2 == 0)
        for cell in ws[row_idx]:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = left_align
            if use_alt:
                cell.fill = alt_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

    ws.freeze_panes = "A2"

    try:
        wb.save(str(output_path))
        return output_path
    except PermissionError:
        alt_path = output_path.parent / f"{output_path.stem}_new{output_path.suffix}"
        wb.save(str(alt_path))
        print(f"\n[NOTE] '{output_path.name}' is currently open in Excel.")
        print(f"       Saved to '{alt_path.name}' instead. Close Excel to overwrite.")
        return alt_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract clean Company, Name, Phone, Email from MNC sheet to Excel."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_XLSX)
    args = parser.parse_args()

    rows = read_mnc_rows(args.input.resolve())
    contacts = clean_contacts(rows, "91")
    saved_path = save_to_excel(args.output.resolve(), contacts)

    print("=" * 60)
    print("           MNC CLEANING COMPLETED")
    print("=" * 60)
    print(f"Total Contact Records Extracted: {len(contacts)}")
    print(f"Clean Excel File Saved:          {saved_path.resolve()}")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
